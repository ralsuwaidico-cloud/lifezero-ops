#!/usr/bin/env python3
"""Check the live article's worked example against the workbook a buyer actually downloads.

growth/pages/reseller-fees-2026.html publishes a five-marketplace table for one $45 item, and
ends by telling the reader "I built a free spreadsheet that does exactly this table for your
item." That sentence is a promise the article and the calculator agree. Nothing was checking it.

data/facts.json (2026-09-20) pinned the RATES this storefront quotes. It cannot catch a wrong
SUM: the article could quote every rate correctly and still print a profit figure that the
calculator contradicts, and a reader who downloaded the file would find the difference before I
did.

The comparison is deliberately NOT a second Python implementation of the fee maths. Re-deriving
it here would repeat the mistake of 2026-09-14, when a build assertion was written from the same
wrong Facebook rate as the code it was meant to guard and so agreed with the bug. Instead the
article is compared against the CACHED, LIBREOFFICE-RECALCULATED values in the built workbook -
the actual numbers a buyer sees when they open the file. The calculator's default item (cost $12,
list $45, no buyer-paid shipping received, $6.50 label, $0.50 other) is the article's item, so
the two are directly comparable with nothing restated in between.

Exit 0 = the article agrees with the workbook. Non-zero = the difference, per cell.
Run: python3 scripts/verify_article_math.py
"""
import decimal
import pathlib
import re
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
ARTICLE = ROOT / "growth" / "pages" / "reseller-fees-2026.html"
BOOK = ROOT / "build" / "out" / "Reseller_Profit_After_Fees_Calculator_2026.xlsx"
ROWS = range(12, 18)          # the marketplace rows on the Calculator sheet


def shown(value, places):
    """What a spreadsheet DISPLAYS for this number, not what Python's round() returns.

    Found by this checker on its first run, on the one row where it matters. Depop leaves
    $43.065; the article prints $43.07 and Python's round() gives 43.06, because 43.065 is
    stored as 43.064999999999998 and binary rounding goes down. Excel and Google Sheets round
    half away from zero on the decimal value, so a buyer opening the file sees $43.07 and the
    article is right. A float tolerance cannot settle this - a half-cent tolerance sits exactly
    on the boundary these values land on, which is why the first version passed three cells at
    .xx5 by luck and failed the fourth. Compare what is DISPLAYED, exactly.
    """
    q = decimal.Decimal(10) ** -places
    return decimal.Decimal(repr(value)).quantize(q, rounding=decimal.ROUND_HALF_UP)


def money(s):
    return float(s.replace("$", "").replace(",", "").strip())


def article_table():
    """The second table in the article: Marketplace | Fee | You receive | Profit | ROI."""
    html = ARTICLE.read_text()
    tables = re.findall(r"<table>.*?</table>", html, re.S)
    if len(tables) < 2:
        return None
    rows = {}
    for tr in re.findall(r"<tr>(.*?)</tr>", tables[1], re.S):
        cells = [re.sub(r"<[^>]+>", "", c).strip()
                 for c in re.findall(r"<t[dh]>(.*?)</t[dh]>", tr, re.S)]
        if len(cells) != 5 or cells[0] == "Marketplace":
            continue
        rows[cells[0]] = {"fee": money(cells[1]), "receive": money(cells[2]),
                          "profit": money(cells[3]),
                          "roi": float(cells[4].replace("%", "")) / 100}
    return rows


def workbook_table():
    from openpyxl import load_workbook
    ws = load_workbook(BOOK, data_only=True)["Calculator"]
    out = {}
    for r in ROWS:
        name = ws.cell(r, 1).value
        fee = ws.cell(r, 2).value
        if not name or fee is None:
            continue
        out[name] = {"fee": fee, "receive": ws.cell(r, 4).value,
                     "profit": ws.cell(r, 6).value, "roi": ws.cell(r, 7).value}
    return out


def main():
    if not ARTICLE.exists():
        print("verify_article_math: no article, nothing to check")
        return 0
    if not BOOK.exists():
        print("verify_article_math: the calculator has not been built; run build/build_all.py "
              "first", file=sys.stderr)
        return 1
    art, book = article_table(), workbook_table()
    if not art:
        print("WRONG: could not find the worked-example table in the article - if it was "
              "removed or restructured, this check is now verifying nothing")
        return 1

    problems = []
    for name, a in sorted(art.items()):
        b = book.get(name)
        if b is None:
            problems.append(f"{name}: in the article but not a row in the calculator - one of "
                            f"the two has gained or lost a marketplace")
            continue
        for field, places, pct in (("fee", 2, False), ("receive", 2, False),
                                   ("profit", 2, False), ("roi", 3, True)):
            want, got = shown(a[field], places), shown(b[field], places)
            if want != got:
                unit = (lambda d: f"{d * 100:.1f}%") if pct else (lambda d: f"${d}")
                problems.append(f"{name} {field}: the article says {unit(want)}, the workbook a "
                                f"buyer downloads shows {unit(got)}")

    # The headline claim is a subtraction the table does not show, so check it separately.
    if art and not problems:
        spread = max(v["profit"] for v in book.values()) - min(v["profit"] for v in book.values())
        claimed = re.search(r"\$(\d+\.\d\d) of profit separates the best from the worst",
                            ARTICLE.read_text())
        if not claimed:
            problems.append("the 'X of profit separates the best from the worst' sentence is "
                            "gone or reworded - the headline number is no longer checked")
        elif shown(float(claimed.group(1)), 2) != shown(spread, 2):
            problems.append(f"headline spread: the article claims ${float(claimed.group(1)):.2f} "
                            f"between best and worst, the workbook computes ${spread:.2f}")

    for p in problems:
        print("WRONG:", p)
    if not problems:
        print(f"verify_article_math: {len(art)} marketplace row(s) in the article match the "
              f"built calculator, and so does the headline spread")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
