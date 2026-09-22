#!/usr/bin/env python3
"""Check the UAE tracker's VAT-return and Corporate-Tax arithmetic, including the branch the
shipped sample data never exercises.

These are the highest-consequence numbers this storefront produces. The reseller calculator gets
a pricing decision wrong; the VAT Return sheet is copied into EmaraTax and filed with a tax
authority, and the Corporate Tax sheet tells someone what to set aside.

build_all already cross-checks that every income row lands in some VAT box and every expense has
a recognised treatment. Those are COMPLETENESS checks. Nothing checked the ARITHMETIC: that Box 8
is the sum of its parts, that net VAT is output minus input, that taxable turnover excludes
exempt and out-of-scope supplies, or that Corporate Tax is 9% of the profit above AED 375,000.

The invariants below are written from the STRUCTURE of the return - net = output - input, and so
on - not by re-deriving the sheet's own formulas. That distinction is the whole point: on
2026-09-14 a build assertion was written from the same wrong rate as the code it guarded and
agreed with the bug. An invariant that comes from the tax form rather than from the spreadsheet
cannot agree with the spreadsheet's mistake.

And the part that matters most: **with the shipped sample data, Small Business Relief applies and
Corporate Tax is zero, so the 9% branch never runs.** Reading zero off clean data proves nothing
about the formula that charges tax. This script therefore builds a second scenario - one invoice
raised to AED 4,000,000, which pushes revenue past the AED 3,000,000 relief threshold - recalculates
it with LibreOffice, and checks that relief is correctly refused and 9% correctly charged. That is
the branch that was silently wrong until 2026-09-14, when the relief window was capped three years
early.

Exit 0 = every invariant holds, in both scenarios. Non-zero = listed.
Run: python3 scripts/verify_uae_returns.py
"""
import glob
import json
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile

ROOT = pathlib.Path(__file__).resolve().parent.parent
BOOK = ROOT / "build" / "out" / "UAE_Business_Bookkeeping_VAT_CT_Tracker_2026.xlsx"
STD_RATE = 0.05
CT_RATE = 0.09
CT_BAND = 375_000
CENT = 0.005


def recalc_script():
    for c in (glob.glob(os.path.expanduser("~/.claude/skills/**/xlsx/scripts/recalc.py"), recursive=True)
              + glob.glob("/root/.claude/skills/**/xlsx/scripts/recalc.py", recursive=True)):
        return pathlib.Path(c)
    return None


def read(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    v, c, s = wb["VAT Return"], wb["Corporate Tax"], wb["Settings"]
    return {
        "std_amt": v.cell(9, 3).value, "std_vat": v.cell(9, 4).value,
        "zero_amt": v.cell(10, 3).value, "zero_vat": v.cell(10, 4).value,
        "exempt_amt": v.cell(11, 3).value, "exempt_vat": v.cell(11, 4).value,
        "out_amt": v.cell(12, 3).value,
        "box8_amt": v.cell(13, 3).value, "box8_vat": v.cell(13, 4).value,
        "box9_amt": v.cell(14, 3).value, "box9_vat": v.cell(14, 4).value,
        "box11_vat": v.cell(15, 4).value, "net_vat": v.cell(16, 4).value,
        "turnover": v.cell(20, 4).value,
        "revenue": c.cell(5, 2).value, "deductible": c.cell(6, 2).value,
        "irrec_vat": c.cell(7, 2).value, "adjust": c.cell(8, 2).value,
        "profit": c.cell(9, 2).value, "elected": c.cell(10, 2).value,
        "eligible": c.cell(11, 2).value, "above_band": c.cell(12, 2).value,
        "ct": c.cell(13, 2).value, "eff_rate": c.cell(14, 2).value,
        "monthly": c.cell(15, 2).value,
        "sbr_threshold": s.cell(13, 2).value,
    }


def close(a, b):
    return a is not None and b is not None and abs(a - b) <= CENT


def invariants(d, label):
    p = []

    def check(ok, msg):
        if not ok:
            p.append(f"{label}: {msg}")

    # --- the VAT return, box by box, from the structure of the form ---
    check(close(d["std_vat"], d["std_amt"] * STD_RATE),
          f"Box 1a output VAT is {d['std_vat']}, but 5% of {d['std_amt']} is "
          f"{d['std_amt'] * STD_RATE}")
    check(close(d["zero_vat"], 0), f"zero-rated supplies carry VAT of {d['zero_vat']}, must be 0")
    check(close(d["exempt_vat"], 0), f"exempt supplies carry VAT of {d['exempt_vat']}, must be 0")
    check(close(d["box8_amt"], d["std_amt"] + d["zero_amt"] + d["exempt_amt"]),
          f"Box 8 total is {d['box8_amt']}, but 1a + 4 + 5 is "
          f"{d['std_amt'] + d['zero_amt'] + d['exempt_amt']} - out-of-scope income must be "
          f"excluded from the return")
    check(close(d["box8_vat"], d["std_vat"] + d["zero_vat"] + d["exempt_vat"]),
          f"Box 8 output VAT is {d['box8_vat']}, but the boxes above sum to "
          f"{d['std_vat'] + d['zero_vat'] + d['exempt_vat']}")
    check(close(d["box9_vat"], d["box9_amt"] * STD_RATE),
          f"Box 9 input VAT is {d['box9_vat']}, but 5% of {d['box9_amt']} is "
          f"{d['box9_amt'] * STD_RATE}")
    check(close(d["box11_vat"], d["box9_vat"]),
          f"Box 11 recoverable input VAT is {d['box11_vat']}, Box 9 is {d['box9_vat']}")
    check(close(d["net_vat"], d["box8_vat"] - d["box11_vat"]),
          f"net VAT payable is {d['net_vat']}, but output {d['box8_vat']} minus input "
          f"{d['box11_vat']} is {d['box8_vat'] - d['box11_vat']} - this is the figure the buyer "
          f"pays the FTA")
    # Taxable turnover for registration is standard + zero-rated only.
    check(close(d["turnover"], d["std_amt"] + d["zero_amt"]),
          f"registration turnover is {d['turnover']}, but standard + zero-rated is "
          f"{d['std_amt'] + d['zero_amt']} - exempt and out-of-scope must be excluded")

    # --- Corporate Tax ---
    want_profit = d["revenue"] - d["deductible"] - d["irrec_vat"] + d["adjust"]
    check(close(d["profit"], want_profit),
          f"taxable profit is {d['profit']}, but revenue - deductible - irrecoverable VAT + "
          f"adjustments is {want_profit}")
    relief = (d["elected"] == "Yes") and d["revenue"] <= d["sbr_threshold"]
    check((d["eligible"] == "Yes") == relief,
          f"Small Business Relief eligibility reads {d['eligible']!r}, but elected="
          f"{d['elected']!r} with revenue {d['revenue']} against a threshold of "
          f"{d['sbr_threshold']} means it should be {'Yes' if relief else 'No'}")
    want_band = max(0.0, d["profit"] - CT_BAND)
    check(close(d["above_band"], want_band),
          f"taxable income above the band is {d['above_band']}, but profit {d['profit']} minus "
          f"AED {CT_BAND:,} is {want_band}")
    want_ct = 0.0 if relief else want_band * CT_RATE
    check(close(d["ct"], want_ct),
          f"estimated Corporate Tax is {d['ct']}, but "
          f"{'relief applies so it must be 0' if relief else f'9% of {want_band} is {want_ct}'}")
    check(close(d["monthly"], want_ct / 12),
          f"suggested monthly set-aside is {d['monthly']}, but CT {want_ct} over 12 months is "
          f"{want_ct / 12}")
    if d["profit"]:
        check(close(d["eff_rate"], want_ct / d["profit"]),
              f"effective rate is {d['eff_rate']}, but {want_ct} on a profit of {d['profit']} is "
              f"{want_ct / d['profit']}")
    return p


def scenario_over_threshold(problems):
    """Raise one invoice past the relief threshold and prove the 9% branch actually charges."""
    rc = recalc_script()
    if rc is None:
        problems.append("scenario: the xlsx skill's recalc.py was not found, so the 9% Corporate "
                        "Tax branch is UNTESTED - it never runs on the shipped sample data")
        return
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as td:
        copy = pathlib.Path(td) / "over_threshold.xlsx"
        shutil.copy(BOOK, copy)
        wb = load_workbook(copy)           # formulas, not values
        wb["Income"].cell(4, 6).value = 4_000_000
        wb.save(copy)
        r = subprocess.run([sys.executable, str(rc), str(copy), "180"],
                           capture_output=True, text=True)
        try:
            info = json.loads(r.stdout)
        except json.JSONDecodeError:
            problems.append(f"scenario: could not recalculate the over-threshold workbook "
                            f"({r.stderr.strip()[:200]})")
            return
        if info.get("status") != "success" or info.get("total_errors"):
            problems.append(f"scenario: the over-threshold workbook recalculated with errors: "
                            f"{info}")
            return
        d = read(copy)
        problems.extend(invariants(d, "over-threshold scenario"))
        # And say plainly that the branch fired, rather than trusting it did.
        if d["eligible"] == "Yes":
            problems.append("over-threshold scenario: revenue of AED 4,012,500 is above the AED "
                            "3,000,000 relief threshold but the sheet still reports relief as "
                            "available - this is the 2026-09-14 class of bug")
        elif not d["ct"]:
            problems.append("over-threshold scenario: relief is correctly refused but Corporate "
                            "Tax still computes 0 - the 9% branch is not charging")
        else:
            print(f"  scenario ok  revenue AED {d['revenue']:,.0f} exceeds the relief threshold: "
                  f"relief refused, CT AED {d['ct']:,.2f} on AED {d['above_band']:,.2f} above the "
                  f"band")


def main():
    if not BOOK.exists():
        print("verify_uae_returns: the UAE tracker has not been built; run build/build_all.py "
              "first", file=sys.stderr)
        return 1
    problems = invariants(read(BOOK), "shipped sample data")
    scenario_over_threshold(problems)
    for p in problems:
        print("WRONG:", p)
    if not problems:
        print("verify_uae_returns: VAT boxes and Corporate Tax reconcile on the shipped data, "
              "and the 9% branch charges correctly when relief is refused")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
