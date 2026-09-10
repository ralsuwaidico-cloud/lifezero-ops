"""Build script for the FREE Reseller Profit-After-Fees Calculator 2026.

The lead magnet for the $19 tracker. It has to be genuinely useful on its own — a crippled
free tool earns a bad first rating, and Gumroad Discover needs a rating before it will show
the paid product at all. So this does one job completely: type an item once, see the profit
on every marketplace side by side, and see which one actually pays most.

Fee rates are the same table the paid tracker uses, with the same per-row "Rates last
checked" provenance. Rows that were not re-verified say so, in red.

Run: python3 build/build_free_calc.py   (openpyxl only; recalculate with recalc.py after).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = "/home/user/lifezero-ops/build/out/Reseller_Profit_After_Fees_Calculator_2026.xlsx"

UPSELL = "https://gum.co/u/ysjnz0ld"   # UTM: free_calculator / upsell / reseller-lite-2026

ACCENT = "1F3A5F"
TINT = "E9EEF5"
BLUE = "0000FF"
RED = "C00000"
FONT = "Arial"

f = lambda **k: Font(name=FONT, **k)
hdr_font = f(bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=ACCENT)
tint_fill = PatternFill("solid", fgColor=TINT)
thin = Side(style="thin", color="BFC7D1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
CUR = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%'

wb = Workbook()


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f(bold=True, size=16, color=ACCENT)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(italic=True, size=9, color="666666")


def header(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = hdr_font; c.fill = hdr_fill; c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 34


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


# --------------------------------------------------------------- FEES
# Columns: name, fee %, per-order $, per-order $ on small orders, small-order limit $,
#          processing %, fee charged on buyer-paid shipping?, flat fee under $, flat fee $,
#          minimum fee $, last checked, note
FEES = [
    ("eBay", 0.136, 0.40, 0.30, 10.0, 0.0, "Y", 0.0, 0.0, 0.0, "2026-09-10",
     "Final value fee 13.6% of the total (item + buyer-paid shipping + handling) plus $0.40 "
     "per order, or $0.30 on orders of $10 or less. Most categories; books/media and some "
     "others differ."),
    ("Poshmark", 0.20, 0.0, 0.0, 0.0, 0.0, "N", 15.0, 2.95, 0.0, "2026-09-10",
     "Flat $2.95 on sales under $15, otherwise 20% of the sale price. No separate processing "
     "fee. The buyer pays shipping directly, so leave 'buyer-paid shipping you receive' at 0."),
    ("Mercari", 0.10, 0.0, 0.0, 0.0, 0.0, "Y", 0.0, 0.0, 0.0, "2026-09-10",
     "Flat 10% on item + buyer-paid shipping. The old 2.9% + $0.50 seller processing fee was "
     "removed on 6 Jan 2025; buyers now pay a 3.6% Buyer Protection fee instead."),
    ("Depop", 0.0, 0.45, 0.45, 0.0, 0.033, "Y", 0.0, 0.0, 0.0, "not re-verified",
     "US selling fee moved to the buyer; roughly 3.3% + $0.45 payment processing remains on "
     "the seller. Verify for your region before trusting this row."),
    ("Facebook Marketplace", 0.05, 0.0, 0.0, 0.0, 0.0, "Y", 0.0, 0.0, 0.40, "not re-verified",
     "About 5% of the total on shipped orders, with a minimum around $0.40. Local pickup is "
     "normally free. Verify before trusting this row."),
    ("Other / edit me", 0.10, 0.0, 0.0, 0.0, 0.0, "Y", 0.0, 0.0, 0.0, "n/a",
     "Blank slot — put any other marketplace's rates here."),
]
NAMES = [p[0] for p in FEES]
FEE_FIRST = 4
FEE_LAST = FEE_FIRST + len(FEES) - 1

fs = wb.active
fs.title = "Fees"
title(fs, "Marketplace fee rates",
      "Blue cells are yours to edit. Every calculation on the Calculator sheet reads from here.")
header(fs, 3, ["Marketplace", "Fee %", "Per-order fee $", "Per-order fee $ (small orders)",
               "Small order up to $", "Processing %", "Fee charged on buyer-paid shipping? (Y/N)",
               "Flat fee applies under $", "Flat fee $", "Minimum fee $", "Rates last checked",
               "Notes — check against the marketplace's own fee page"])
for i, p in enumerate(FEES):
    r = FEE_FIRST + i
    for c, v in enumerate(p, start=1):
        cell = fs.cell(row=r, column=c, value=v)
        cell.border = border
        cell.font = f(size=10, color=BLUE) if c in (2, 3, 4, 5, 6, 7, 8, 9, 10) else f(size=10)
        if c in (2, 6):
            cell.number_format = PCT
        if c in (3, 4, 5, 8, 9, 10):
            cell.number_format = CUR
        if c == 1:
            cell.font = f(size=10, bold=True)
        if c == 11:
            # An unverified rate that looks verified is worse than one that admits it.
            cell.font = f(size=10, bold=True, color=RED if v == "not re-verified" else "006100")
            cell.alignment = Alignment(horizontal="center")
        if c == 12:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
widths(fs, [22, 9, 13, 15, 14, 12, 16, 15, 11, 13, 15, 78])
fs.freeze_panes = "A4"

# --------------------------------------------------------------- CALCULATOR
cs = wb.create_sheet("Calculator")
wb.move_sheet("Calculator", offset=-1)
title(cs, "Reseller profit-after-fees calculator 2026",
      "Type your item into the blue cells once. Every marketplace is priced below, side by side.")

inputs = [
    ("What you paid for the item", 12.00, CUR, "Your cost of goods — the thrift/wholesale price."),
    ("Sale price you plan to list at", 45.00, CUR, "The item price the buyer pays, before shipping."),
    ("Buyer-paid shipping YOU receive", 0.00, CUR,
     "0 if you offer free shipping, or if the marketplace collects it and buys the label "
     "(Poshmark). Put the amount here only when it lands in your payout (typical on eBay)."),
    ("Your shipping cost (label + packaging)", 6.50, CUR,
     "What the label and box/poly-mailer actually cost you. Enter it even when shipping is free "
     "to the buyer — it comes out of your pocket either way."),
    ("Other costs on this item", 0.50, CUR, "Cleaning, repair, gas share, supplies."),
]
r = 4
for label, val, fmt, note in inputs:
    cs.cell(row=r, column=1, value=label).font = f(bold=True, size=10)
    c = cs.cell(row=r, column=2, value=val)
    c.font = f(size=11, color=BLUE, bold=True); c.number_format = fmt
    c.fill = tint_fill; c.border = border
    n = cs.cell(row=r, column=3, value=note)
    n.font = f(size=9, color="666666"); n.alignment = Alignment(wrap_text=True, vertical="top")
    # The notes are long; give them room instead of letting them spill over the table below.
    cs.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    cs.row_dimensions[r].height = 30
    r += 1

COST, PRICE, SHIP_IN, SHIP_OUT, OTHER = "$B$4", "$B$5", "$B$6", "$B$7", "$B$8"
SPEND = f"({COST}+{SHIP_OUT}+{OTHER})"

cs["A10"] = "What each marketplace leaves you"
cs["A10"].font = f(bold=True, size=12, color=ACCENT)
header(cs, 11, ["Marketplace", "Fee", "Fee as % of what the buyer pays", "You receive",
                "Your total outlay", "PROFIT", "ROI", "Rates last checked"])

CROW = 12
for i in range(len(FEES)):
    r = CROW + i
    fr = FEE_FIRST + i                       # matching row on the Fees sheet
    base = f"({PRICE}+IF(Fees!$G${fr}=\"Y\",{SHIP_IN},0))"
    per_order = f"IF(AND(Fees!$E${fr}>0,{PRICE}<=Fees!$E${fr}),Fees!$D${fr},Fees!$C${fr})"
    tiered = (f"IF(AND(Fees!$H${fr}>0,{PRICE}<Fees!$H${fr}),Fees!$I${fr},"
              f"{base}*Fees!$B${fr}+{per_order}+{base}*Fees!$F${fr})")

    cs.cell(row=r, column=1, value=f"=Fees!$A${fr}").font = f(bold=True, size=10)
    cs.cell(row=r, column=2, value=f"=MAX({tiered},Fees!$J${fr})").number_format = CUR
    # Guarding the division: a $0 sale price is a perfectly normal thing to type while
    # experimenting, and #DIV/0! across the row makes the sheet look broken.
    cs.cell(row=r, column=3, value=f'=IF(({PRICE}+{SHIP_IN})<=0,"-",B{r}/({PRICE}+{SHIP_IN}))').number_format = PCT
    cs.cell(row=r, column=4, value=f"={PRICE}+{SHIP_IN}-B{r}").number_format = CUR
    cs.cell(row=r, column=5, value=f"={SPEND}").number_format = CUR
    cs.cell(row=r, column=6, value=f"=D{r}-E{r}").number_format = CUR
    cs.cell(row=r, column=7, value=f'=IF({SPEND}<=0,"-",F{r}/{SPEND})').number_format = PCT
    cs.cell(row=r, column=8, value=f"=Fees!$K${fr}").alignment = Alignment(horizontal="center")
    for c in range(1, 9):
        cs.cell(row=r, column=c).border = border
    cs.cell(row=r, column=6).font = f(bold=True, size=11)

LAST = CROW + len(FEES) - 1
BEST = LAST + 2
cs.cell(row=BEST, column=1, value="Best marketplace for this item").font = f(bold=True, size=12, color=ACCENT)
cs.cell(row=BEST, column=2,
        value=f"=INDEX(A{CROW}:A{LAST},MATCH(MAX(F{CROW}:F{LAST}),F{CROW}:F{LAST},0))").font = \
    f(bold=True, size=12, color="006100")
cs.cell(row=BEST + 1, column=1, value="Profit there").font = f(bold=True, size=10)
c = cs.cell(row=BEST + 1, column=2, value=f"=MAX(F{CROW}:F{LAST})")
c.number_format = CUR; c.font = f(bold=True, size=11, color="006100")
cs.cell(row=BEST + 2, column=1, value="Worst on this list costs you").font = f(bold=True, size=10)
c = cs.cell(row=BEST + 2, column=2, value=f"=MAX(F{CROW}:F{LAST})-MIN(F{CROW}:F{LAST})")
c.number_format = CUR; c.font = f(bold=True, size=11, color=RED)
cs.cell(row=BEST + 2, column=3,
        value="…more than the best one, on this single item. That is the whole point of "
              "checking before you list.").font = f(size=9, italic=True, color="666666")

# Loss warning. Google Sheets rejects a direct cross-sheet reference inside a conditional
# formatting formula, but the range here is on this sheet, so a plain relative formula is fine.
cs.conditional_formatting.add(
    f"A{CROW}:H{LAST}",
    FormulaRule(formula=[f"$F{CROW}<0"], fill=PatternFill("solid", fgColor="FFC7CE"),
                font=Font(name=FONT, color="9C0006")))

note_row = BEST + 4
cs.cell(row=note_row, column=1,
        value="Promoted listings, store subscriptions, international/ad fees and sales tax are "
              "not modelled. Rates marked 'not re-verified' are approximations — open the "
              "Fees sheet and check them against the marketplace's own fee page.")\
    .font = f(size=9, italic=True, color="666666")
cs.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)

cs.cell(row=note_row + 2, column=1,
        value="Tracking a whole business, not one item?").font = f(bold=True, size=11, color=ACCENT)
cs.cell(row=note_row + 3, column=1,
        value="This sheet prices one item at a time. The full Reseller Inventory & Profit "
              "Tracker 2026 does the rest of the job: 500 items of inventory with auto SKUs, "
              "every sale's fees and ROI calculated as you type, a death-pile report for "
              "anything that has sat too long, sell-through and best-platform dashboards, and "
              "a Schedule-C-style tax summary at year end. $19, one file, no subscription — "
              f"{UPSELL}").alignment = Alignment(wrap_text=True, vertical="top")
cs.cell(row=note_row + 3, column=1).font = f(size=10)
cs.merge_cells(start_row=note_row + 3, start_column=1, end_row=note_row + 3, end_column=8)
cs.row_dimensions[note_row + 3].height = 58

widths(cs, [34, 16, 16, 14, 15, 14, 10, 16])
cs.sheet_view.showGridLines = False

# --------------------------------------------------------------- BATCH
bs = wb.create_sheet("Batch (20 items)")
title(bs, "Price 20 items at once",
      "Pick a marketplace per row. Same maths as the Calculator sheet, one row per item.")
header(bs, 3, ["Item", "Marketplace", "Item cost $", "Sale price $", "Buyer-paid shipping you receive $",
               "Your shipping cost $", "Other costs $", "Fee", "You receive", "PROFIT", "ROI"])
BF, BL = 4, 23
dv = DataValidation(type="list", formula1='"' + ",".join(NAMES) + '"', allow_blank=True)
bs.add_data_validation(dv)
dv.add(f"B{BF}:B{BL}")

for r in range(BF, BL + 1):
    # MATCH the chosen marketplace back to its Fees row, then repeat the fee maths.
    m = f"MATCH($B{r},Fees!$A${FEE_FIRST}:$A${FEE_LAST},0)"
    g = f"INDEX(Fees!$G${FEE_FIRST}:$G${FEE_LAST},{m})"
    base = f"($D{r}+IF({g}=\"Y\",$E{r},0))"
    idx = lambda col: f"INDEX(Fees!${col}${FEE_FIRST}:${col}${FEE_LAST},{m})"
    per_order = f"IF(AND({idx('E')}>0,$D{r}<={idx('E')}),{idx('D')},{idx('C')})"
    tiered = (f"IF(AND({idx('H')}>0,$D{r}<{idx('H')}),{idx('I')},"
              f"{base}*{idx('B')}+{per_order}+{base}*{idx('F')})")
    spend = f"($C{r}+$F{r}+$G{r})"

    for c in range(1, 8):
        cell = bs.cell(row=r, column=c)
        cell.font = f(size=10, color=BLUE); cell.border = border
        if c >= 3:
            cell.number_format = CUR
    # An empty row must stay empty rather than reporting a fee on nothing.
    bs.cell(row=r, column=8, value=f'=IF($B{r}="","",MAX({tiered},{idx("J")}))').number_format = CUR
    bs.cell(row=r, column=9, value=f'=IF($B{r}="","",$D{r}+$E{r}-$H{r})').number_format = CUR
    bs.cell(row=r, column=10, value=f'=IF($B{r}="","",$I{r}-{spend})').number_format = CUR
    bs.cell(row=r, column=11, value=f'=IF(OR($B{r}="",{spend}<=0),"",$J{r}/{spend})').number_format = PCT
    bs.cell(row=r, column=10).font = f(bold=True, size=10)
    for c in range(8, 12):
        bs.cell(row=r, column=c).border = border

bs.cell(row=BF, column=1, value="Example: Patagonia fleece")
bs.cell(row=BF, column=2, value="eBay")
for col, v in zip("CDEFG", [12, 45, 0, 6.5, 0.5]):
    bs[f"{col}{BF}"] = v

bs.cell(row=BL + 2, column=1, value="TOTAL").font = f(bold=True, size=11, color=ACCENT)
c = bs.cell(row=BL + 2, column=10, value=f"=SUM(J{BF}:J{BL})")
c.number_format = CUR; c.font = f(bold=True, size=11)

bs.conditional_formatting.add(
    f"A{BF}:K{BL}",
    FormulaRule(formula=[f'AND($B{BF}<>"",$J{BF}<0)'], fill=PatternFill("solid", fgColor="FFC7CE"),
                font=Font(name=FONT, color="9C0006")))
widths(bs, [26, 20, 12, 12, 16, 15, 12, 12, 13, 13, 10])
bs.freeze_panes = "A4"

# --------------------------------------------------------------- START HERE
ss = wb.create_sheet("START HERE")
wb.move_sheet("START HERE", offset=-(len(wb.sheetnames) - 1))
title(ss, "Reseller profit-after-fees calculator 2026",
      "Free. eBay · Poshmark · Mercari · Depop · Facebook Marketplace")
rows = [
    ("WHAT THIS DOES", None),
    ("One item, every platform",
     "Type an item's cost, price and shipping into the Calculator sheet once. It shows the fee, "
     "what you receive, your profit and your ROI on every marketplace at the same time, and "
     "names the one that pays most. A row turns red when you would lose money."),
    ("20 at a time",
     "The 'Batch (20 items)' sheet does the same maths one row per item, with a dropdown to pick "
     "the marketplace per row."),
    ("", ""),
    ("BEFORE YOU TRUST A NUMBER", None),
    ("Check the rates",
     "Marketplaces change fees without much notice, and a stale rate quietly distorts every "
     "figure here. The Fees sheet carries a 'Rates last checked' date per row. eBay, Poshmark "
     "and Mercari were verified on 10 September 2026. Depop and Facebook Marketplace say 'not "
     "re-verified' in red — that is the honest state of those two rows. Five minutes on each "
     "platform's fee page and you can fix them yourself; everything recalculates."),
    ("Buyer-paid shipping",
     "Only enter shipping that actually lands in your payout. On Poshmark the buyer pays the "
     "label directly, so leave it at 0. On eBay it usually does reach you — and eBay charges "
     "its percentage on it, which this sheet accounts for."),
    ("Not modelled",
     "Promoted listings, store subscriptions, international and advertising fees, and sales tax. "
     "Add anything like that to 'Other costs'."),
    ("", ""),
    ("HOW TO OPEN IT", None),
    ("Excel", "Excel 2016 or newer. No macros, nothing to enable."),
    ("Google Sheets", "File → Import → Upload, then 'Replace spreadsheet'. Everything works."),
    ("", ""),
    ("IF THIS IS USEFUL", None),
    ("Rate it",
     "A rating on the Gumroad page is the single most useful thing you can give back — it is "
     "what lets anyone else find this. It costs you ten seconds."),
    ("The full tracker",
     "This prices one item. The Reseller Inventory & Profit Tracker 2026 runs the business: 500 "
     "inventory items with auto-generated SKUs, fees and ROI on every sale as you type, a death "
     "pile report for stale listings with suggested markdowns, dashboards for sell-through and "
     "best platform, and a Schedule-C-style tax summary at year end. $19, one file, no "
     f"subscription. {UPSELL}"),
    ("", ""),
    ("DISCLAIMER", None),
    ("Not tax or financial advice",
     "Fee rates are user-editable approximations. Verify them for your own account and region "
     "before making pricing decisions."),
]
r = 4
for a, b in rows:
    if b is None:
        c = ss.cell(row=r, column=1, value=a)
        c.font = f(bold=True, color="FFFFFF", size=11); c.fill = hdr_fill
        ss.cell(row=r, column=2).fill = hdr_fill
    else:
        ca = ss.cell(row=r, column=1, value=a); ca.font = f(bold=True, size=10, color=ACCENT)
        ca.alignment = Alignment(vertical="top")
        cb = ss.cell(row=r, column=2, value=b); cb.font = f(size=10)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
widths(ss, [26, 108])
ss.sheet_view.showGridLines = False

wb.save(OUT)
print("saved", OUT)
