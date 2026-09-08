"""Build: Project Brief & Intake Pack — the file a $95 custom-sheet buyer receives at purchase.

Gumroad requires a product to deliver real content on purchase; a listing that only tells the
buyer to email is rejected ("This product has no content attached and directs buyers to message
you on another platform"). This pack is that content, and it earns its place: a structured brief
gets the build right first time instead of costing a round-trip out of the 48 hours.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "/home/claude/products/custom-sheet-48h/Custom_Sheet_Project_Brief.xlsx"

FONT = "Arial"
ACCENT = "1F3A5F"
LIGHT = "EAF1FB"
BLUE = "0000FF"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=ACCENT)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(italic=True, color="595959")
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------- HOW THIS WORKS
ws = wb.active
ws.title = "START HERE"
title(ws, "Your custom spreadsheet — how this works",
      "You are holding step 1. Fill in the Brief tab, send it back, and the 48 hours begins.")
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 112

rows = [
    ("THE THREE STEPS", True),
    ("1. Fill in the 'Brief' tab in this workbook. It takes about five minutes. If you have sample data, paste it into the 'Data Sample' tab.", False),
    ("2. Send this file back as a reply to your Gumroad receipt email.", False),
    ("3. You get a first version within 48 hours of that reply landing, then one revision round to get it right.", False),
    ("", False),
    ("THE 48 HOURS STARTS WHEN YOUR BRIEF ARRIVES", True),
    ("Not at checkout. Nothing can be built before the brief exists, so there is no clock running while you think about it. Take the time to fill it in properly — a clear brief is the difference between one revision and three.", False),
    ("", False),
    ("WHAT YOU GET", True),
    ("A working spreadsheet built to your brief: formulas checked, dropdowns where a list is expected, a dashboard that updates itself, and a one-page how-to written for your sheet. It opens in Microsoft Excel (2016 or later) and in Google Sheets (File → Import). No macros, nothing to install.", False),
    ("", False),
    ("WHAT IS NOT INCLUDED", True),
    ("Macros or VBA. Connections to a live database or API. Anything that files tax for you. Ongoing data entry or maintenance. If your brief turns out to need one of these, you will be told straight away and refunded — not handed something half-built.", False),
    ("", False),
    ("BEFORE YOU START", True),
    ("The single most useful thing you can give is a sample of your real data, even five rows, with the column headings you actually use. Second most useful: a sentence about the decision you want to make from the sheet. 'Which products are losing me money' produces a better tool than 'a sales tracker'.", False),
]
r = 4
for text, bold in rows:
    c = ws.cell(row=r, column=2, value=text)
    c.font = f(bold=bold, color=ACCENT if bold else "000000", size=11 if bold else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if bold:
        c.fill = fill(LIGHT)
    r += 1

# ---------------------------------------------------------------- BRIEF
br = wb.create_sheet("Brief")
title(br, "Your brief", "Type in the blue cells. Nothing here is compulsory except questions 1 and 2.")
br.column_dimensions["A"].width = 46
br.column_dimensions["B"].width = 60
br.column_dimensions["C"].width = 54

questions = [
    ("1. What should the sheet track?", "", "One or two sentences is plenty."),
    ("2. What decision do you want to make from it?", "", "e.g. 'which clients are slow payers', 'what to reorder'."),
    ("3. Excel or Google Sheets?", "Excel", "Pick from the dropdown. Both are supported."),
    ("4. Roughly how many rows do you add per month?", "", "A number, or 'no idea'. It decides how the sheet is sized."),
    ("5. What should the output look like?", "", "A dashboard? A monthly summary? One number at the top? Describe it, or paste a screenshot into the Data Sample tab."),
    ("6. Currency / units", "", "e.g. USD, AED, GBP, kg."),
    ("7. Anything that must NOT change", "", "e.g. 'keep my existing column order', 'my accountant needs the current layout'."),
    ("8. Deadline, if you have one", "", "Beyond the 48 hours — say so if a specific date matters."),
    ("9. Anything else", "", "Quirks, exceptions, the thing that makes your business different."),
]
hdr = ["Question", "Your answer", "Notes"]
for i, h in enumerate(hdr, 1):
    c = br.cell(row=3, column=i, value=h)
    c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    c.fill = fill(ACCENT)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = BORDER
br.row_dimensions[3].height = 24

r = 4
for q, default, note in questions:
    a = br.cell(row=r, column=1, value=q)
    a.font = f(bold=True)
    a.alignment = Alignment(wrap_text=True, vertical="top")
    b = br.cell(row=r, column=2, value=default)
    b.font = f(color=BLUE)
    b.alignment = Alignment(wrap_text=True, vertical="top")
    n = br.cell(row=r, column=3, value=note)
    n.font = f(italic=True, color="595959")
    n.alignment = Alignment(wrap_text=True, vertical="top")
    for col in (1, 2, 3):
        br.cell(row=r, column=col).border = BORDER
    br.row_dimensions[r].height = 42
    r += 1

dv = DataValidation(type="list", formula1='"Excel,Google Sheets,Either"', allow_blank=True)
br.add_data_validation(dv)
dv.add("B6")

br.cell(row=r + 1, column=1, value="When this is filled in, reply to your Gumroad receipt email with this file attached.").font = f(bold=True, color=ACCENT, size=11)

# ---------------------------------------------------------------- DATA SAMPLE
ds = wb.create_sheet("Data Sample")
title(ds, "Data sample (optional, but the most useful thing you can send)",
      "Paste a few real rows below — headings included. Five rows is enough. Redact anything sensitive.")
for i in range(1, 13):
    c = ds.cell(row=4, column=i, value=f"Column {i}")
    c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    c.fill = fill(ACCENT)
    c.border = BORDER
    ds.column_dimensions[ds.cell(row=4, column=i).column_letter].width = 18
for row in range(5, 26):
    for col in range(1, 13):
        ds.cell(row=row, column=col).border = BORDER
ds.cell(row=28, column=1, value="Replace the headings above with your own. If your data lives somewhere else, a screenshot pasted onto this tab works just as well.").font = f(italic=True, color="595959")
ds.sheet_view.showGridLines = False

for ws_, col in zip(wb.worksheets, [ACCENT, BLUE, "808080"]):
    ws_.sheet_properties.tabColor = col
wb.active = 0
wb.save(OUT)
print("saved", OUT)
