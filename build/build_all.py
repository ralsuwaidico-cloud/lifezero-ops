#!/usr/bin/env python3
"""Regenerate every product asset from source (text-only kit -> binaries).
Run from anywhere: python3 build/build_all.py
Produces products/<slug>/*.xlsx and cover *.jpg, recalculated and verified."""
import glob, json, os, pathlib, re, subprocess, sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
PROD = ROOT / "products"
WORK = ROOT / "build" / "out"
WORK.mkdir(parents=True, exist_ok=True)

# LibreOffice needs well over 120s on a cold container; override with RECALC_TIMEOUT.
RECALC_TIMEOUT = os.environ.get("RECALC_TIMEOUT", "600")

# locate the xlsx skill's recalc + soffice helper (synced skills dir)
skill = None
for c in glob.glob(os.path.expanduser("~/.claude/skills/**/xlsx/scripts/recalc.py"), recursive=True) + \
         glob.glob("/root/.claude/skills/**/xlsx/scripts/recalc.py", recursive=True):
    skill = pathlib.Path(c).parent; break
if not skill:
    sys.exit("xlsx skill scripts not found")

def build(script, out_name):
    src = (ROOT / "build" / script).read_text()
    target = WORK / out_name
    src = re.sub(r'^OUT\s*=\s*.*$', f'OUT = "{target}"', src, count=1, flags=re.M)
    tmp = WORK / f"_{script}"
    tmp.write_text(src)
    subprocess.run([sys.executable, str(tmp)], check=True)
    r = subprocess.run([sys.executable, str(skill / "recalc.py"), str(target), RECALC_TIMEOUT], capture_output=True, text=True)
    info = json.loads(r.stdout)
    assert info.get("status") == "success" and info.get("total_errors") == 0, info
    print("built", target.name, info["total_formulas"], "formulas, 0 errors")
    return target

uae = build("build_uae.py", "UAE_Business_Bookkeeping_VAT_CT_Tracker_2026.xlsx")
res = build("build_reseller.py", "Reseller_Inventory_Profit_Tracker_2026.xlsx")


def cross_check(path, a, b, label):
    """Assert two cells that the workbook itself promises are equal really are.

    'Formulas recalculate with 0 errors' does not mean the numbers are right: the reseller
    Dashboard once computed Net Profit from the shipping column and still recalculated
    cleanly, reporting a loss on a profitable year. Only comparing the totals catches that.
    """
    from openpyxl import load_workbook as _lw
    wb = _lw(path, data_only=True)
    va = wb[a[0]][a[1]].value or 0
    vb = wb[b[0]][b[1]].value or 0
    assert abs(va - vb) < 0.005, f"{label}: {a[0]}!{a[1]}={va} != {b[0]}!{b[1]}={vb}"
    print(f"cross-check ok  {label}: {va}")


cross_check(res, ("Tax Summary", "C21"), ("Dashboard", "J24"),
            "reseller net profit (Tax Summary line 15 vs Dashboard total)")


def assert_zero(path, sheet, cell, label):
    from openpyxl import load_workbook as _lw
    v = _lw(path, data_only=True)[sheet][cell].value or 0
    assert abs(v) < 0.005, f"{label}: {sheet}!{cell}={v}, expected 0"
    print(f"cross-check ok  {label}: 0")


# Verified to fire: setting Income!G4 to "Standard" instead of "Standard 5%" drops that invoice
# from every VAT box and takes net VAT payable from 630 to 30, with 0 formula errors throughout.
assert_zero(uae, "VAT Return", "D17", "UAE income fully classified into VAT boxes")
assert_zero(uae, "VAT Return", "D19", "UAE expenses all have a recognised VAT treatment")

# preview render of dashboards -> PNG
from openpyxl import load_workbook
def preview(path, keep):
    wb = load_workbook(path, data_only=True)
    for ws in list(wb.worksheets):
        if ws.title not in keep: wb.remove(ws)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    p = WORK / (path.stem + "_preview.xlsx"); wb.save(p)
    subprocess.run([sys.executable, str(skill / "office" / "soffice.py"), "--headless", "--convert-to", "pdf", "--outdir", str(WORK), str(p)],
                   capture_output=True)
    pdf = WORK / (p.stem + ".pdf")
    subprocess.run(["pdftoppm", "-png", "-r", "110", "-f", "1", "-l", "1", str(pdf), str(WORK / p.stem)], check=True)
    return next(WORK.glob(p.stem + "*-1.png"))

from PIL import Image, ImageDraw, ImageFont
F_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"; F_R = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
def cover(png, title, subtitle, out, bbox):
    im = Image.open(png).convert("RGB").crop(bbox)
    W, H = 1280, 720
    scale = min((W - 80) / im.width, (H - 170) / im.height)
    im = im.resize((int(im.width * scale), int(im.height * scale)), Image.LANCZOS)
    bg = Image.new("RGB", (W, H), (31, 58, 95)); d = ImageDraw.Draw(bg)
    d.text((40, 30), title, fill="white", font=ImageFont.truetype(F_B, 40))
    d.text((40, 88), subtitle, fill=(200, 215, 235), font=ImageFont.truetype(F_R, 22))
    x = (W - im.width) // 2; y = 140
    d.rectangle([x - 6, y - 6, x + im.width + 6, y + im.height + 6], fill="white"); bg.paste(im, (x, y))
    bg.save(out, quality=92); print("cover", out.name)


def _trim_whitespace(im, pad=14):
    """LibreOffice renders a whole page, so a short sheet leaves most of the frame blank.
    Crop to the actual content before framing it."""
    from PIL import ImageChops
    grey = im.convert("L").point(lambda v: 0 if v > 247 else 255)
    box = grey.getbbox()
    if not box:
        return im
    l, t, r, b = box
    return im.crop((max(0, l - pad), max(0, t - pad),
                    min(im.width, r + pad), min(im.height, b + pad)))


def page_shot(png, title, out):
    """A gallery image of one whole sheet. Buyers cannot open a spreadsheet before paying, so
    the listing should show them the actual sheets, not just one dashboard."""
    im = _trim_whitespace(Image.open(png).convert("RGB"))
    W, H = 1280, 720
    scale = min((W - 60) / im.width, (H - 110) / im.height)
    im = im.resize((max(1, int(im.width * scale)), max(1, int(im.height * scale))), Image.LANCZOS)
    bg = Image.new("RGB", (W, H), (31, 58, 95))
    d = ImageDraw.Draw(bg)
    d.text((40, 26), title, fill="white", font=ImageFont.truetype(F_B, 34))
    x = (W - im.width) // 2; y = 92
    d.rectangle([x - 5, y - 5, x + im.width + 5, y + im.height + 5], fill="white")
    bg.paste(im, (x, y))
    bg.save(out, quality=92)
    print("preview", out.name)


def square_thumb(png, title, out, bbox=None):
    """Gumroad rejects a non-square thumbnail ("Please upload a square thumbnail"),
    so the 16:9 cover cannot double as one. Build a real 1000x1000 tile."""
    S = 1000
    bg = Image.new("RGB", (S, S), (31, 58, 95))
    d = ImageDraw.Draw(bg)
    font = ImageFont.truetype(F_B, 62)
    # wrap the title to the tile width
    words, lines, cur = title.split(), [], ""
    for w in words:
        trial = f"{cur} {w}".strip()
        if d.textlength(trial, font=font) <= S - 100:
            cur = trial
        else:
            lines.append(cur); cur = w
    if cur:
        lines.append(cur)
    y = 60
    for ln in lines:
        d.text((50, y), ln, fill="white", font=font); y += 74
    if png is not None:
        im = Image.open(png).convert("RGB")
        if bbox:
            im = im.crop(bbox)
        avail_h = S - y - 70
        scale = min((S - 100) / im.width, avail_h / im.height)
        im = im.resize((max(1, int(im.width * scale)), max(1, int(im.height * scale))), Image.LANCZOS)
        x = (S - im.width) // 2; ty = y + 30
        d.rectangle([x - 6, ty - 6, x + im.width + 6, ty + im.height + 6], fill="white")
        bg.paste(im, (x, ty))
    bg.save(out, quality=92)
    assert bg.width == bg.height, "thumbnail must be square"
    print("thumb", out.name, f"{bg.width}x{bg.height}")

cover(preview(uae, ["Dashboard"]), "UAE Freelancer & SME Bookkeeping Tracker 2026",
      "VAT return helper  •  Corporate Tax estimate  •  Invoices, expenses & cash — Excel / Google Sheets",
      PROD / "uae-vat-tracker" / "cover_uae_tracker.jpg", (70, 95, 910, 810))
cover(preview(res, ["Dashboard"]), "Reseller Inventory & Profit Tracker 2026",
      "eBay • Poshmark • Mercari • Depop • Whatnot • FB Marketplace — auto fees, ROI, death pile, tax summary",
      PROD / "reseller-profit-tracker" / "cover_reseller_tracker.jpg", (70, 95, 800, 810))

# service cover (pure text)
bg = Image.new("RGB", (1280, 720), (31, 58, 95)); d = ImageDraw.Draw(bg)
d.text((70, 120), "Custom Excel / Google Sheets Tool", fill="white", font=ImageFont.truetype(F_B, 54))
d.text((70, 190), "built for you in 48 hours", fill="white", font=ImageFont.truetype(F_B, 54))
for i, t in enumerate(["Trackers  •  Calculators  •  Dashboards  •  Data clean-up", "Formulas checked. Dropdowns. One-page how-to.", "One revision round included."]):
    d.text((70, 300 + i * 50), t, fill=(200, 215, 235), font=ImageFont.truetype(F_R, 30))
d.rounded_rectangle([70, 500, 470, 570], radius=12, fill=(46, 117, 182)); d.text((100, 515), "Brief in → sheet out", fill="white", font=ImageFont.truetype(F_B, 30))
bg.save(PROD / "custom-sheet-48h" / "cover_custom_sheet.jpg", quality=92)

# Gallery previews - what the buyer actually gets, sheet by sheet.
page_shot(preview(uae, ["VAT Return"]), "VAT Return helper — copy these boxes into EmaraTax",
          PROD / "uae-vat-tracker" / "preview_uae_vat_return.jpg")
page_shot(preview(uae, ["Income"]), "Income — one row per invoice, VAT calculated for you",
          PROD / "uae-vat-tracker" / "preview_uae_income.jpg")
page_shot(preview(res, ["Sales"]), "Sales — fees, profit and ROI calculated per sale",
          PROD / "reseller-profit-tracker" / "preview_reseller_sales.jpg")
page_shot(preview(res, ["Tax Summary"]), "Tax Summary — Schedule-C-style totals for the year",
          PROD / "reseller-profit-tracker" / "preview_reseller_tax.jpg")

# Intake pack for the $95 service. Gumroad requires a product to deliver real content on
# purchase - a listing that only says "email me" is rejected as having nothing attached.
brief = build("build_brief.py", "Custom_Sheet_Project_Brief.xlsx")

# Square thumbnails - Gumroad rejects the 16:9 cover for this slot.
_uae_png = preview(uae, ["Dashboard"])
_res_png = preview(res, ["Dashboard"])
square_thumb(_uae_png, "UAE VAT & Corporate Tax Tracker 2026",
             PROD / "uae-vat-tracker" / "thumb_uae_tracker.jpg", (70, 95, 910, 810))
square_thumb(_res_png, "Reseller Inventory & Profit Tracker 2026",
             PROD / "reseller-profit-tracker" / "thumb_reseller_tracker.jpg", (70, 95, 800, 810))
square_thumb(None, "Custom Excel / Google Sheets Tool in 48 Hours",
             PROD / "custom-sheet-48h" / "thumb_custom_sheet.jpg")

import shutil
shutil.copy(uae, PROD / "uae-vat-tracker" / uae.name)
shutil.copy(res, PROD / "reseller-profit-tracker" / res.name)
shutil.copy(brief, PROD / "custom-sheet-48h" / brief.name)
print("ALL ASSETS BUILT")
